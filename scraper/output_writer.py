from __future__ import annotations

import csv
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .normalize import social_label

COLUMNS = ["Nama", "Der Treffpunkt", "Sosmed", "Kontak", "Gambar", "Alasan"]
CSV_NAME = "hasil_scraping.csv"
XLSX_NAME = "hasil_scraping.xlsx"
ZIP_NAME = "hasil_scraping.zip"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_LINK_FONT = Font(color="0563C1", underline="single")


def _csv_safe(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_safe(str(row.get(key, ""))) for key in COLUMNS})


def _link_label(column: str, url: str, name: str) -> str:
    if column == "Der Treffpunkt":
        return f"Maps {name}"
    if column == "Sosmed":
        return social_label(url, name)
    if column == "Gambar":
        return "galeri foto"
    return url


def _write_xlsx(rows: list[dict], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hasil"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="563CF5")

    for row_index, row in enumerate(rows, start=2):
        name = row.get("Nama", "")
        for col_index, column in enumerate(COLUMNS, start=1):
            value = str(row.get(column, ""))
            cell = sheet.cell(row=row_index, column=col_index)
            if value and column in {"Der Treffpunkt", "Sosmed", "Kontak", "Gambar"}:
                cell.value = _link_label(column, value, name)
                cell.hyperlink = value
                cell.font = _LINK_FONT
            else:
                cell.value = value
            cell.data_type = "s"

    for column, width in zip("ABCDEF", (34, 30, 30, 30, 14, 50)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def _write_zip(output_dir: Path, path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in (CSV_NAME, XLSX_NAME):
            archive.write(output_dir / name, name)
        images_dir = output_dir / "images"
        if images_dir.is_dir():
            for image in sorted(images_dir.iterdir()):
                archive.write(image, f"images/{image.name}")


def write_outputs(rows: list[dict], output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {"csv": output_dir / CSV_NAME, "xlsx": output_dir / XLSX_NAME, "zip": output_dir / ZIP_NAME}
    _write_csv(rows, paths["csv"])
    _write_xlsx(rows, paths["xlsx"])
    _write_zip(output_dir, paths["zip"])
    return paths
